from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


IMAGE = "Curriculo_Brendon_da_Silva_Pinto_Maurilio.png"
OUTPUT = "Curriculo_Brendon_da_Silva_Pinto_Maurilio.docx"


document = Document()
section = document.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(0)
section.bottom_margin = Cm(0)
section.left_margin = Cm(0)
section.right_margin = Cm(0)
section.header_distance = Cm(0)
section.footer_distance = Cm(0)

paragraph = document.add_paragraph()
paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph.paragraph_format.space_before = Pt(0)
paragraph.paragraph_format.space_after = Pt(0)
paragraph.paragraph_format.line_spacing = Pt(1)

run = paragraph.add_run()
inline_shape = run.add_picture(IMAGE, width=Cm(20.99))

# Add accessibility metadata while keeping the page artwork unchanged.
doc_pr = inline_shape._inline.docPr
doc_pr.set("name", "Currículo de Brendon da Silva Pinto Maurilio")
doc_pr.set(
    "descr",
    "Currículo profissional para a área de produção, sem foto, em página A4.",
)

# Prevent Word from adding compatibility spacing around the full-page artwork.
settings = document.settings._element
compat = settings.find(qn("w:compat"))
if compat is None:
    compat = OxmlElement("w:compat")
    settings.append(compat)
do_not_use_html_auto_spacing = OxmlElement("w:doNotUseHTMLParagraphAutoSpacing")
compat.append(do_not_use_html_auto_spacing)

document.save(OUTPUT)
